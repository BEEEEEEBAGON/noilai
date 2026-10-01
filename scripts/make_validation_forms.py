#!/usr/bin/env python
"""Build the native-validation packet and the human-baseline forms; score what comes back.

Protocol: docs/gate1/VALIDATION_PROTOCOL.md (DESIGN_DECISIONS 10.1 / 10.2 as amended on 1 October 2026,
docs/DEVIATIONS.md). Sizes come from noilai.constants (chosen with scripts/validation_sizing.py).

    # the packet: Parts A-E for each validator (one .xlsx workbook each, CSV copies beside it) + the author's keys
    python scripts/make_validation_forms.py packet --release data/release/v0.3 --out data/validation \
        --validators A B C --attested data/attested_seed.tsv --candidates data/attested_candidates.tsv

    # after the main round: third-validator sheet for the split rows (three validators only)
    python scripts/make_validation_forms.py adjudication-sheet --dir data/validation --to C

    # score everything returned (validation_<V>.xlsx or <sheet>_<V>.csv files under --returned)
    python scripts/make_validation_forms.py score --dir data/validation --returned 'data/validation/returned/*'

    # human baseline (DESIGN_DECISIONS 10.2): 20 forms x 30 items with the models' exact p0 prompt, + check item,
    # + the 10-item natural-competence block; then score returned forms with the models' scorer
    python scripts/make_validation_forms.py baseline --items data/release/v0.3/noilai_main.jsonl --out data/human \
        --attested data/release/v0.3/attested.jsonl
    python scripts/make_validation_forms.py score-baseline --items data/release/v0.3/noilai_main.jsonl --dir data/human \
        --returned 'data/human/returned/*.csv'

Validators are letters; nothing identifying is stored; data/validation/ and data/human/ are git-ignored and never
leave the author's machine except as the sheets sent to the validators (DESIGN_DECISIONS 11.2). Seeds are PUBLIC
sampling seeds, never the withheld build seed (DESIGN_DECISIONS 4.6).
"""
from __future__ import annotations

import argparse
import csv
import datetime as _dt
import glob
import json
import random
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from noilai import constants as C
from noilai import validation as VA
from noilai.gen.generate import load_items

PACKET_SEED = 20261018               # public sampling seed of the validation packet (Gate 1, 18 October 2026)
BASELINE_SEED = 20261102            # public sampling seed (docs/HUMAN_BASELINE_FORM.md); never the build seed (design 4.6)
BASELINE_CELLS = [(t, v) for t in ("T1", "T2", "T3") for v in ("V1", "V2", "V3", "V4")]
ANCHOR_CELLS = [("T1", "V1"), ("T1", "V2"), ("T1", "V3"), ("T1", "V4"), ("T2", "V1"), ("T3", "V1")]   # design 10.2: 6 anchors
NATURAL_BLOCK_ITEMS = 10            # design 10.2: T2 on attested, non-vulgar rows, same for every respondent
CHECK_ITEM = {"item_id": "CHECK-01", "task": "check", "variant": "",
              "prompt_vi": "Câu kiểm tra: xin hãy viết đúng hai chữ “đã đọc” vào ô trả lời của câu này. [NATIVE-CHECK]",
              "expected": "đã đọc"}
# closing questions (PREREG 5.7): "yes" to either excludes the form; asked on every form, answered Có / Không
CLOSING_ITEMS = (
    {"item_id": "TOOLS", "prompt_vi": "Bạn có dùng từ điển, công cụ tìm kiếm hay trợ lý AI cho câu nào không? (Có/Không) "
                                      "Xin trả lời thật; câu trả lời “Có” không ảnh hưởng gì đến bạn. [NATIVE-CHECK]"},
    {"item_id": "WAS_VALIDATOR", "prompt_vi": "Bạn có phải là người kiểm định (đánh giá bảng dữ liệu) của dự án NóiLái không? "
                                              "(Có/Không) [NATIVE-CHECK]"},
)

SHEETS = {
    "A_calibration": ["row_id", "task", "kind", "input", "candidate", "system_verdict", "question_vi", "question_en",
                      "correct", "spelling", "lexical_input", "lexical_candidate", "offensive", "dialect", "comment"],
    "B_items": ["row_id", "task", "kind", "input", "candidate", "system_verdict", "question_vi", "question_en",
                "correct", "spelling", "lexical_input", "lexical_candidate", "offensive", "dialect", "comment"],
    "C_attested": ["row_id", "input", "output", "source", "known", "valid", "spelling_ok", "your_form", "offensive",
                   "dialect", "comment"],
    "D1_production": ["row_id", "phrase", "your_noilai", "comment"],
    "D2_qu": ["row_id", "phrase", "kind", "option_A", "option_B", "choice", "your_form", "comment"],
    "D3_conventions": ["row_id", "topic", "question_vi", "options", "answer", "comment"],
    "E_t2gold": ["row_id", "noilai_form", "reading_1", "accept_1", "reading_2", "accept_2", "reading_3", "accept_3",
                 "missing_reading", "comment"],
}
YESNO_COLUMNS = {"correct", "spelling", "lexical_input", "lexical_candidate", "offensive", "known", "valid", "spelling_ok",
                 "accept_1", "accept_2", "accept_3"}


# --------------------------------------------------------------------------- io
def _write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in fields})


def _write_xlsx(path: Path, sheets: dict[str, list[dict]]) -> bool:
    """One workbook per validator, a dropdown on every yes/no and dialect column (imports into Google Sheets with
    the dropdowns). Returns False when openpyxl is not installed (the CSV copies are then the sheets to send)."""
    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        fields = SHEETS[name]
        ws = wb.create_sheet(name)
        ws.append(fields)
        for r in rows:
            ws.append([r.get(k, "") for k in fields])
        n = max(2, len(rows) + 1)
        yn = DataValidation(type="list", formula1='"Có,Không,Không chắc"', allow_blank=True)
        dia = DataValidation(type="list", formula1='"' + ",".join(VA.DIALECT_CHOICES) + '"', allow_blank=True)
        ab = DataValidation(type="list", formula1='"A,B,cả hai,không cái nào"', allow_blank=True)
        ws.add_data_validation(yn)
        ws.add_data_validation(dia)
        ws.add_data_validation(ab)
        for j, col in enumerate(fields, start=1):
            letter = ws.cell(row=1, column=j).column_letter
            rng = f"{letter}2:{letter}{n}"
            if col in YESNO_COLUMNS:
                yn.add(rng)
            elif col == "dialect":
                dia.add(rng)
            elif col == "choice":
                ab.add(rng)
        ws.freeze_panes = "B2"
    wb.save(path)
    return True


def _read_returned(patterns: list[str]) -> dict[str, dict[str, list[dict]]]:
    """{sheet: {validator: rows}} from validation_<V>.xlsx workbooks and/or <sheet>_<V>.csv files. Rows are keyed by
    row_id per validator, so a CSV copy of a workbook sheet does not double-count, and a returned third-validator
    sheet (B_adjudicate_<V>.csv) adds rows to that validator's B_items."""
    acc: dict = defaultdict(lambda: defaultdict(dict))
    paths = sorted({p for pat in patterns for p in glob.glob(pat)})
    for p in paths:
        path = Path(p)
        if path.suffix == ".xlsx":
            from openpyxl import load_workbook
            v = path.stem.split("_")[-1]
            wb = load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                if ws.title not in SHEETS:
                    continue
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue
                head = [str(h) if h is not None else "" for h in rows[0]]
                for r in rows[1:]:
                    if any(c not in (None, "") for c in r):
                        d = {h: ("" if c is None else str(c)) for h, c in zip(head, r)}
                        acc[ws.title][v][d["row_id"]] = d
        elif path.suffix == ".csv":
            stem = path.stem
            sheet, v = stem.rsplit("_", 1)
            if sheet not in SHEETS and sheet != "B_adjudicate":
                continue
            with open(path, encoding="utf-8") as fh:
                rows = list(csv.DictReader(fh))
            for d in rows:
                acc["B_items" if sheet == "B_adjudicate" else sheet][v][d["row_id"]] = d
    return {sheet: {v: list(rows.values()) for v, rows in by_v.items()} for sheet, by_v in acc.items()}


def _sha(path: Path) -> str | None:
    if not path.exists():
        return None
    from noilai.eval.run import sha256_file
    return sha256_file(path)


def _read_attested_rows(seed_tsv: Path | None, candidates_tsv: Path | None) -> list[dict]:
    rows = []
    for path, tag in ((seed_tsv, "seed"), (candidates_tsv, "candidate")):
        if path is None or not Path(path).exists():
            continue
        with open(path, encoding="utf-8") as fh:
            for r in csv.DictReader(fh, delimiter="\t"):
                if not (r.get("input") and r.get("output")):
                    continue
                if tag == "candidate" and r.get("status", "candidate") in ("rejected", "duplicate_of_seed"):
                    continue
                src = r.get("source") or r.get("title") or ""
                if tag == "candidate" and r.get("url"):
                    src = f"{r.get('title', '')} <{r['url']}>".strip()
                rows.append({"input": r["input"], "output": r["output"], "source": src, "origin": tag})
    return rows


# --------------------------------------------------------------------------- packet
def cmd_packet(args) -> int:
    rel = Path(args.release)
    test = load_items(rel / "noilai_test.jsonl")
    dev = load_items(rel / "noilai_dev.jsonl")
    core = load_items(rel / "noilai_core.jsonl")
    items = test + dev
    vals = args.validators
    if len(vals) < 2:
        raise SystemExit("native validation needs at least two validators")
    keys = VA.release_phrase_keys(items)
    calib = VA.build_calibration(keys)
    cal_keys = {VA.phrase_key(r["base_phrase"]) for r in calib} | {VA.phrase_key(r["candidate"]) for r in calib}
    part_b = VA.build_part_b(items, vals, args.seed, args.per_cell, args.controls_per_cell, args.overlap, cal_keys)
    att_rows = _read_attested_rows(Path(args.attested) if args.attested else None,
                                   Path(args.candidates) if args.candidates else None)
    part_c = VA.build_part_c(att_rows)
    sup = VA.build_supplementary()
    part_e = VA.build_part_e(core, vals, args.seed + 1, overlap=args.t2_overlap) if not args.no_t2_gold else None
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    rows_b = {r["row_id"]: r for r in part_b["rows"]}
    rows_e = {r["row_id"]: r for r in part_e["rows"]} if part_e else {}
    per_validator = {}
    wrote_xlsx = False
    for v in vals:
        sheets = {
            "A_calibration": [{k: r[k] for k in r if k not in ("key", "explanation_vi", "explanation_en", "manipulation", "base_phrase")}
                              for r in calib],
            "B_items": [rows_b[rid] for rid in part_b["sheets"][v]],
            "C_attested": part_c,
            "D1_production": sup["D1"], "D2_qu": sup["D2"], "D3_conventions": sup["D3"],
        }
        if part_e:
            sheets["E_t2gold"] = [rows_e[rid] for rid in part_e["sheets"].get(v, [])]
        for name, rows in sheets.items():
            _write_csv(out / f"{name}_{v}.csv", SHEETS[name], rows)
        wrote_xlsx |= _write_xlsx(out / f"validation_{v}.xlsx", sheets)
        n_b, n_c, n_e = len(sheets["B_items"]), len(part_c), len(sheets.get("E_t2gold", []))
        hours = (len(calib) * C.VALIDATION_SECONDS_PER_ITEM + n_b * C.VALIDATION_SECONDS_PER_ITEM + n_c * 25
                 + (len(sup["D1"]) + len(sup["D2"]) + len(sup["D3"])) * 30 + n_e * 20) / 3600
        per_validator[v] = {"A": len(calib), "B": n_b, "C": n_c, "D": len(sup["D1"]) + len(sup["D2"]) + len(sup["D3"]),
                            "E": n_e, "hours_estimate": round(hours, 1)}
    # the author's keys (never sent to a validator)
    (out / "A_calibration_key.json").write_text(json.dumps(calib, ensure_ascii=False, indent=1), encoding="utf-8")
    (out / "B_key.json").write_text(json.dumps(part_b["key"], ensure_ascii=False, indent=1), encoding="utf-8")
    (out / "C_rows.json").write_text(json.dumps(part_c, ensure_ascii=False, indent=1), encoding="utf-8")
    (out / "D_engine.json").write_text(json.dumps(sup, ensure_ascii=False, indent=1), encoding="utf-8")
    if part_e:
        (out / "E_key.json").write_text(json.dumps(part_e["key"], ensure_ascii=False, indent=1), encoding="utf-8")
    manifest = {
        "created_utc": _dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "protocol": "docs/gate1/VALIDATION_PROTOCOL.md", "seed": args.seed, "validators": vals,
        "release": str(rel), "release_files_sha256": {n: _sha(rel / n) for n in ("noilai_test.jsonl", "noilai_dev.jsonl", "noilai_core.jsonl")},
        "attested_sources": {"seed": args.attested, "candidates": args.candidates,
                             "seed_sha256": _sha(Path(args.attested)) if args.attested else None,
                             "candidates_sha256": _sha(Path(args.candidates)) if args.candidates else None},
        "sizes": {"per_cell": args.per_cell, "controls_per_cell": args.controls_per_cell, "overlap": args.overlap,
                  "n_sample": part_b["n_sample"], "n_controls": part_b["n_controls"], "calibration": len(calib),
                  "attested_rows": len(part_c), "D1": len(sup["D1"]), "D2": len(sup["D2"]), "D3": len(sup["D3"]),
                  "t2_gold_items": len(part_e["rows"]) if part_e else 0},
        "strata": part_b["strata"], "per_validator": per_validator, "xlsx_written": wrote_xlsx,
        "keys_never_sent": ["A_calibration_key.json", "B_key.json", "D_engine.json", "E_key.json"],
    }
    (out / "validation_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: manifest[k] for k in ("sizes", "per_validator", "xlsx_written")}, ensure_ascii=False))
    return 0


# --------------------------------------------------------------------------- adjudication sheet
def cmd_adjudication_sheet(args) -> int:
    d = Path(args.dir)
    key = json.loads((d / "B_key.json").read_text(encoding="utf-8"))
    returned = _read_returned(args.returned)
    rep = VA.score_part_b(returned.get("B_items", {}), key, n_boot=0 if args.fast else 200)
    rows_b = {}
    sheet_src = {}
    for path in sorted(d.glob("B_items_*.csv")):
        with open(path, encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                sheet_src[r["row_id"]] = r
    for a in rep["adjudication"]["rows"]:
        if a["rule"] != "needs_author" or args.to in a["judgments"]:
            continue
        r = dict(sheet_src[a["row_id"]])
        for j in VA.JUDGMENTS_B + ("dialect", "comment"):
            r[j] = ""
        rows_b[a["row_id"]] = r
    path = d / f"B_adjudicate_{args.to}.csv"
    _write_csv(path, SHEETS["B_items"], list(rows_b.values()))
    print(json.dumps({"written": str(path), "rows": len(rows_b)}, ensure_ascii=False))
    return 0


# --------------------------------------------------------------------------- score
def _read_author_decisions(path: Path) -> dict:
    if not path.exists():
        return {}
    with open(path, encoding="utf-8") as fh:
        return {r["row_id"]: r["decision"].strip().lower() for r in csv.DictReader(fh, delimiter="\t")
                if r.get("decision", "").strip().lower() in ("yes", "no")}


def cmd_score(args) -> int:
    d = Path(args.dir)
    returned = _read_returned(args.returned)
    regions = {}
    if (d / "validators.json").exists():
        regions = {k: v.get("region") for k, v in json.loads((d / "validators.json").read_text(encoding="utf-8")).items()}
    rep_dir = d / "report"
    rep_dir.mkdir(parents=True, exist_ok=True)
    report: dict = {"created_utc": _dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "regions": regions}
    # Part A
    if (d / "A_calibration_key.json").exists() and "A_calibration" in returned:
        calib = json.loads((d / "A_calibration_key.json").read_text(encoding="utf-8"))
        report["calibration"] = {v: VA.score_calibration(rows, calib) for v, rows in returned["A_calibration"].items()}
    # Part B
    if (d / "B_key.json").exists() and "B_items" in returned:
        key = json.loads((d / "B_key.json").read_text(encoding="utf-8"))
        decisions = _read_author_decisions(d / "adjudication.tsv")
        b = VA.score_part_b(returned["B_items"], key, decisions, regions, n_boot=args.n_boot)
        report["B"] = {k: b[k] for k in ("n_rows_returned", "n_rows_in_key", "agreement", "agreement_samples_only", "validators",
                                         "generator_precision", "controls_final", "flags")}
        report["B"]["adjudication_counts"] = b["adjudication"]["counts"]
        adj_path = d / "adjudication.tsv"
        old = {}
        if adj_path.exists():
            with open(adj_path, encoding="utf-8") as fh:
                old = {r["row_id"]: r for r in csv.DictReader(fh, delimiter="\t")}
        fields = ["row_id", "item_id", "cell", "control", "judgments", "comments", "decision", "reason", "date"]
        with open(adj_path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=fields, delimiter="\t")
            w.writeheader()
            for a in b["adjudication"]["rows"]:
                prev = old.get(a["row_id"], {})
                w.writerow({"row_id": a["row_id"], "item_id": a["item_id"], "cell": a["cell"], "control": a["control"],
                            "judgments": json.dumps(a["judgments"], ensure_ascii=False),
                            "comments": json.dumps(a["comments"], ensure_ascii=False),
                            "decision": prev.get("decision", ""), "reason": prev.get("reason", ""), "date": prev.get("date", "")})
        with open(rep_dir / "item_flags.tsv", "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow(["row_id", "item_id", "final_correct", "offensive_any", "dialect", "rule_flag"])
            for rid, fl in sorted(b["item_flags"].items()):
                w.writerow([rid, fl["item_id"], b["final"].get(rid) or "", fl["offensive_any"], ",".join(fl["dialect"]), fl["rule_flag"]])
    # Part C
    if (d / "C_rows.json").exists() and "C_attested" in returned:
        rows = json.loads((d / "C_rows.json").read_text(encoding="utf-8"))
        c = VA.score_part_c(returned["C_attested"], rows, n_boot=args.n_boot)
        report["C"] = {k: c[k] for k in ("agreement", "n_rows", "n_verified")}
        with open(rep_dir / "attested_verified.tsv", "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow(["row_id", "input", "output", "source", "verified", "valid_votes", "known_any", "spelling_disputed",
                        "your_forms", "offensive_any", "dialect", "validators"])
            for r in c["rows"]:
                w.writerow([r["row_id"], r["input"], r["output"], r["source"], r["verified"], json.dumps(r["valid"], ensure_ascii=False),
                            r["known_any"], r["spelling_disputed"], " | ".join(r["your_forms"]), r["offensive_any"],
                            ",".join(r["dialect"]), ",".join(r["validators"])])
    # Part D: returned as is, classified by the engine for D1
    if "D1_production" in returned and (d / "D_engine.json").exists():
        eng = {r["row_id"]: r for r in json.loads((d / "D_engine.json").read_text(encoding="utf-8"))["D1"]}
        prod = defaultdict(Counter)
        for v, rows in returned["D1_production"].items():
            for r in rows:
                e = eng.get(r["row_id"])
                ans = VA.canonical_text(r.get("your_noilai", "") or "")
                if not e or not ans:
                    continue
                kind = "other"
                for k, form in e["_engine"].items():
                    fc = VA.canonical_text(form)
                    if ans == fc:
                        kind = k
                    elif set(ans.split()) == set((fc or "").split()):
                        kind = {"V1": "V6", "V2": "V3", "V4": "V5"}[k]
                prod[v][kind] += 1
        report["D1_kind_produced"] = {v: dict(c) for v, c in prod.items()}
    # Part E
    if (d / "E_key.json").exists() and "E_t2gold" in returned:
        key_e = json.loads((d / "E_key.json").read_text(encoding="utf-8"))
        e = VA.score_part_e(returned["E_t2gold"], key_e)
        report["E"] = {k: e[k] for k in ("alpha_masi", "alpha_jaccard", "n_items_returned", "n_items_multi_coded")}
        (rep_dir / "t2_validated_gold.json").write_text(json.dumps(e["validated_gold"], ensure_ascii=False, indent=1), encoding="utf-8")
    (rep_dir / "validation_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print(json.dumps({k: (v if k != "B" else {"precision": v["generator_precision"]["pooled"],
                                              "alpha_correct": v["agreement"]["correct"].get("alpha")})
                      for k, v in report.items() if k in ("B", "C", "E")}, ensure_ascii=False, default=str))
    return 0


# --------------------------------------------------------------------------- human baseline
def baseline_prompt(it: dict) -> str:
    """[NATIVE-CHECK] the short question per task (the wording of docs/HUMAN_BASELINE_FORM.md §1)."""
    if it["task"] == "T1":
        return f"Nói lái kiểu {it['variant']} của “{it['input']}” là gì?"
    if it["task"] == "T2":
        return f"“{it['input']}” là cách nói lái của cụm từ nào?"
    return f"“{it['candidate']}” có phải là nói lái kiểu {it['variant']} của “{it['input']}” không? (Có/Không)"


def model_prompt(it: dict) -> tuple[str, str]:
    """The exact user message the models receive for this item under the main condition (p0, 3 demonstrations,
    explained instruction, raw input, NFC baseline) and its prompt hash."""
    from noilai.eval import prompts as P
    msgs = P.render(it, paraphrase="p0", shots=3, arm="nfc")
    return msgs[-1]["content"], P.prompt_hash(msgs)


def baseline_design(items: list[dict], n_forms: int, per_form: int, seed: int, n_anchors: int = len(ANCHOR_CELLS),
                    per_cell: int | None = None) -> dict:
    """The design 10.2 assignment. Returns {'anchors', 'pool', 'forms' (lists of items, anchors first),
    'per_cell'}; raises ValueError when the sizes do not fit or a cell is short."""
    if n_forms < 2:
        raise ValueError("the double-coverage design needs at least two forms")
    if n_anchors != len(ANCHOR_CELLS):
        raise ValueError(f"design 10.2 fixes {len(ANCHOR_CELLS)} anchors (one per T1 variant, one T2, one T3)")
    slots = n_forms * (per_form - n_anchors)
    if slots <= 0 or slots % 2:
        raise ValueError(f"{n_forms} forms x ({per_form} - {n_anchors} anchors) = {slots} slots must be a positive even number")
    pool_size = slots // 2
    if per_cell is None:
        if pool_size % len(BASELINE_CELLS):
            raise ValueError(f"the {pool_size} double-judged items do not divide over the {len(BASELINE_CELLS)} task x variant cells")
        per_cell = pool_size // len(BASELINE_CELLS)
    elif per_cell * len(BASELINE_CELLS) != pool_size:
        raise ValueError(f"--per-cell {per_cell} x {len(BASELINE_CELLS)} cells != {pool_size} double-judged items")
    rng = random.Random(seed)
    items = [it for it in items if it.get("task") in ("T1", "T2", "T3") and not it.get("vulgar") and it.get("twin_type") != "spelling"]
    by_cell = defaultdict(list)
    for it in items:
        by_cell[(it["task"], it["variant"])].append(it)
    anchors, pool = [], []
    for cell in BASELINE_CELLS:
        c = sorted(by_cell.get(cell, []), key=lambda x: x["item_id"])
        rng.shuffle(c)
        if cell in ANCHOR_CELLS:
            if not c:
                raise ValueError(f"cell {cell} has no eligible item for its anchor")
            anchors.append(c.pop())
            c = [x for x in c if x.get("pair_item_id") != anchors[-1]["item_id"]]     # never the twin of an anchor
        if cell[0] == "T3":                                                            # half yes, half no, never both twins
            n_yes = (per_cell + 1) // 2
            yes = [x for x in c if x["gold"] == "yes"][:n_yes]
            taken = {y["item_id"] for y in yes}
            no = [x for x in c if x["gold"] == "no" and x.get("pair_item_id") not in taken][: per_cell - n_yes]
            picked = yes + no
        else:
            picked = c[:per_cell]
        if len(picked) < per_cell:
            raise ValueError(f"cell {cell} has {len(picked)} eligible items, {per_cell} needed")
        pool += picked
    rng.shuffle(pool)
    forms = [list(anchors) for _ in range(n_forms)]
    for k, it in enumerate(pool):                       # first copy: round robin
        forms[k % n_forms].append(it)
    for k, it in enumerate(pool):                       # second copy: block-specific cyclic offset, never 0
        offset = 1 + (k // n_forms) % (n_forms - 1)
        forms[(k % n_forms + offset) % n_forms].append(it)
    for fm in forms:
        assert len(fm) == per_form and len({x["item_id"] for x in fm}) == per_form
    return {"anchors": anchors, "pool": pool, "forms": forms, "per_cell": per_cell}


def natural_block(attested: list[dict], n: int, seed: int) -> list[dict]:
    """The 10-item natural-competence block (design 10.2): T2 on attested two-syllable, non-vulgar rows."""
    rows = [r for r in attested if not r.get("vulgar") and len((r.get("input") or "").split()) == 2
            and len((r.get("attested_output") or r.get("output") or "").split()) == 2]
    rows = sorted(rows, key=lambda r: (r.get("item_id") or r.get("input")))
    random.Random(seed).shuffle(rows)
    out = []
    for r in rows[:n]:
        lai = r.get("attested_output") or r.get("output")
        out.append({"item_id": r.get("item_id") or f"ATT-{len(out) + 1:02d}", "task": "natural", "variant": "",
                    "prompt_vi": f"“{lai}” là cách nói lái của cụm từ nào? [NATIVE-CHECK]", "expected": r["input"]})
    return out


def cmd_baseline(args) -> int:
    rng = random.Random(args.seed)
    design = baseline_design(load_items(Path(args.items)), args.n_forms, args.per_form, args.seed, args.anchors, args.per_cell)
    anchors, pool, forms = design["anchors"], design["pool"], design["forms"]
    nat = []
    if args.attested and Path(args.attested).exists():
        with open(args.attested, encoding="utf-8") as fh:
            att = [json.loads(line) for line in fh if line.strip()]
        nat = natural_block([r for r in att if "_header" not in r], NATURAL_BLOCK_ITEMS, args.seed)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    fields = ["form", "position", "item_id", "task", "variant", "block", "prompt_model", "prompt_hash", "question_short", "answer"]
    for fi, its in enumerate(forms):
        its = list(its)
        rng.shuffle(its)                                # seeded random order per form (T1, T2, T3 interleaved)
        check_at = rng.randrange(5, len(its))           # the instruction-check item, never among the first five
        rows = []
        for it in its:
            text, h = model_prompt(it) if not args.no_model_prompt else ("", "")
            rows.append({"item_id": it["item_id"], "task": it["task"], "variant": it["variant"], "block": "main",
                         "prompt_model": text, "prompt_hash": h, "question_short": baseline_prompt(it), "answer": ""})
        rows.insert(check_at, {"item_id": CHECK_ITEM["item_id"], "task": "check", "variant": "", "block": "check",
                               "prompt_model": "", "prompt_hash": "", "question_short": CHECK_ITEM["prompt_vi"], "answer": ""})
        for r in nat:
            rows.append({"item_id": r["item_id"], "task": "natural", "variant": "", "block": "natural", "prompt_model": "",
                         "prompt_hash": "", "question_short": r["prompt_vi"], "answer": ""})
        for c in CLOSING_ITEMS:
            rows.append({"item_id": c["item_id"], "task": "closing", "variant": "", "block": "closing", "prompt_model": "",
                         "prompt_hash": "", "question_short": c["prompt_vi"], "answer": ""})
        for pos, r in enumerate(rows, start=1):
            r["form"], r["position"] = fi + 1, pos
        _write_csv(out / f"baseline_form_{fi+1:02d}.csv", fields, rows)
    cover = Counter(it["item_id"] for fm in forms for it in fm)
    anchor_ids = {a["item_id"] for a in anchors}
    seen = [{x["item_id"] for x in fm} - anchor_ids for fm in forms]
    shared = Counter(len(seen[i] & seen[j]) for i in range(len(forms)) for j in range(i + 1, len(forms)))
    info = {"forms": args.n_forms, "per_form": args.per_form, "anchors": len(anchors), "per_cell": design["per_cell"],
            "double_judged": len(pool), "distinct_items": len(cover),
            "min_appearances": min(cover.values()), "max_appearances": max(cover.values()),
            "anchors_seen_by": min(cover[a] for a in anchor_ids),
            "others": sorted({v for k, v in cover.items() if k not in anchor_ids}),
            "shared_items_per_form_pair": {str(k): v for k, v in sorted(shared.items())},
            "rater_graph_connected": 0 not in shared, "seed": args.seed, "items_file": str(args.items),
            "check_items_per_form": 1, "natural_block_items": len(nat), "closing_items": [c["item_id"] for c in CLOSING_ITEMS],
            "design": "DESIGN_DECISIONS 10.2: anchors on every form, every other item on exactly two forms, cyclic double coverage"}
    (out / "human_items.json").write_text(json.dumps(sorted(cover), ensure_ascii=False), encoding="utf-8")
    (out / "baseline_manifest.json").write_text(json.dumps({**info, "anchor_ids": sorted(anchor_ids),
                                                            "appearances": dict(sorted(cover.items())),
                                                            "check_item": CHECK_ITEM, "natural_block": nat},
                                                           ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(info, ensure_ascii=False))
    return 0


def cmd_score_baseline(args) -> int:
    """Score returned baseline forms with the models' scorer (DESIGN_DECISIONS 10.2; PREREG 8.11 and the exclusions
    of PREREG 5.7). A returned CSV keeps the form's columns; `answer` filled; a closing row with item_id `TOOLS`
    and the `WAS_VALIDATOR` row record the closing questions (yes to either excludes the form)."""
    from noilai.eval.score import score_outputs
    from noilai.vi.reencode import canonical_text
    items = {it["item_id"]: it for it in load_items(Path(args.items))}
    man = json.loads((Path(args.dir) / "baseline_manifest.json").read_text(encoding="utf-8"))
    nat_key = {r["item_id"]: r["expected"] for r in man.get("natural_block", [])}
    per_form, excluded, scored = {}, {}, []
    for p in sorted({q for pat in args.returned for q in glob.glob(pat)}):
        with open(p, encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
        form = Path(p).stem
        main = [r for r in rows if r.get("block", "main") == "main" and r["item_id"] in items]
        answered = sum(bool((r.get("answer") or "").strip()) for r in main)
        check = next((r for r in rows if r["item_id"] == CHECK_ITEM["item_id"]), None)
        tools = next((r for r in rows if r["item_id"] == "TOOLS"), None)
        was_validator = next((r for r in rows if r["item_id"] == "WAS_VALIDATOR"), None)
        reason = None
        if answered < args.min_answered:
            reason = f"answered {answered} < {args.min_answered}"
        elif check is not None and canonical_text(check.get("answer") or "") != canonical_text(CHECK_ITEM["expected"]):
            reason = "failed instruction check"
        elif tools is not None and VA.norm_label(tools.get("answer")) == "yes":
            reason = "reported using a dictionary, search engine or AI assistant"
        elif was_validator is not None and VA.norm_label(was_validator.get("answer")) == "yes":
            reason = "respondent was a validator (validators never take the baseline form)"
        if reason:
            excluded[form] = reason
            continue
        outs = [{"item_id": r["item_id"], "task": items[r["item_id"]]["task"], "arm": "base", "prompt_id": f"human-{form}",
                 "raw": r.get("answer") or "", "answer": (r.get("answer") or "").strip() or None, "extraction_method": "human"}
                for r in main]
        rows_scored = score_outputs([items[o["item_id"]] for o in outs], outs)
        for r in rows_scored:
            r["form"] = form
        scored += rows_scored
        nat = [r for r in rows if r.get("block") == "natural"]
        nat_ok = sum(canonical_text(r.get("answer") or "") is not None and
                     set((canonical_text(r.get("answer") or "") or "").split()) == set((canonical_text(nat_key.get(r["item_id"], "")) or "").split())
                     for r in nat)
        per_form[form] = {"n_main": len(main), "answered": answered, "natural_correct": nat_ok, "natural_n": len(nat)}
    by_item = defaultdict(list)
    for r in scored:
        by_item[r["item_id"]].append(bool(r.get("correct")))
    item_acc = {k: sum(v) / len(v) for k, v in by_item.items()}
    by_task = defaultdict(list)
    for k, acc in item_acc.items():
        by_task[items[k]["task"]].append(acc)
    out = Path(args.dir) / "report"
    out.mkdir(parents=True, exist_ok=True)
    with open(out / "human_scores.jsonl", "w", encoding="utf-8") as f:
        for r in scored:
            f.write(json.dumps(r, ensure_ascii=False, default=str) + "\n")
    rep = {"forms_scored": len(per_form), "forms_excluded": excluded, "per_form": per_form,
           "mean_human_by_task": {t: sum(v) / len(v) for t, v in sorted(by_task.items())},
           "n_items_scored": len(item_acc),
           "note": "point estimates only; the two-way (person, item) bootstrap of PREREG 8.11 runs in the analysis step"}
    (out / "human_baseline_report.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(rep, ensure_ascii=False))
    return 0


# --------------------------------------------------------------------------- Google Forms (baseline)
FORM_INTRO_VI = (
    "Cảm ơn bạn đã tham gia. Bạn sẽ làm 30 câu nói lái, giống hệt các câu mà các mô hình ngôn ngữ nhận được (mỗi câu có "
    "phần hướng dẫn và ba ví dụ mẫu; phần hướng dẫn lặp lại ở mỗi câu, bạn chỉ cần đọc kỹ ở vài câu đầu), sau đó 10 câu "
    "giải nói lái quen thuộc và hai câu hỏi cuối. Xin không dùng từ điển, công cụ tìm kiếm hay trợ lý AI; nếu không biết, "
    "hãy để trống hoặc đoán. Thời gian khoảng 30–45 phút. [NATIVE-CHECK]")
CONSENT_BOXES_VI = (
    "Tôi từ 18 tuổi trở lên.",
    "Tôi đã đọc phiếu thông tin và đồng ý tham gia; tôi hiểu việc tham gia là tự nguyện, không được trả tiền và tôi có thể dừng bất cứ lúc nào.",
    "Tôi đồng ý cho lưu các câu trả lời và thông tin nhân khẩu học thô, và cho công bố chúng dưới dạng tổng hợp, ẩn danh.",
)
DEMOGRAPHICS_VI = (
    ("Nhóm tuổi", ("18–29", "30–49", "50 trở lên"), True),
    ("Bạn lớn lên ở vùng nào?", ("Bắc", "Trung", "Nam", "Ngoài Việt Nam"), False),
    ("Bạn đã sống ở Việt Nam bao nhiêu năm?", ("Dưới 5", "5–15", "Trên 15"), False),
    ("Bạn có lớn lên ở nước ngoài không?", ("Có", "Không"), False),
)


def _js(x) -> str:
    return json.dumps(x, ensure_ascii=False)


def cmd_google_form(args) -> int:
    """Write an Apps Script (data/human/build_forms.gs) that creates one Google Form per baseline form in the author's own
    Google account (script.google.com -> paste -> run buildAll). E-mail collection off; consent boxes required; the item
    titles start with the form position so that `import-responses` can map answers back to item ids."""
    d = Path(args.dir)
    forms = []
    for path in sorted(d.glob("baseline_form_*.csv")):
        with open(path, encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
        items = []
        for r in rows:
            yesno = r["task"] in ("T3", "closing")
            items.append({"pos": int(r["position"]), "title": f"{r['position']}. {r['question_short']}",
                          "help": r.get("prompt_model", ""), "yesno": yesno, "required": r["block"] == "closing"})
        forms.append({"form": int(rows[0]["form"]), "items": items})
    if not forms:
        raise SystemExit(f"no baseline_form_*.csv in {d}: run the baseline subcommand first")
    js = f"""// Generated by scripts/make_validation_forms.py google-form ({len(forms)} forms). Do not edit by hand: re-generate.
// Use: script.google.com -> New project (in the author's own Google account) -> paste this file -> run buildAll ->
// authorize -> the log lists each form number with its link. Send each respondent ONE link. E-mail collection is off.
// Afterwards download each form's responses as CSV (Responses -> Sheets -> File -> Download -> CSV), name them
// responses_form_<nn>.csv and run: python scripts/make_validation_forms.py import-responses --dir {d.as_posix()}
const FORMS = {_js(forms)};
const INTRO = {_js(FORM_INTRO_VI)};
const CONSENT = {_js(list(CONSENT_BOXES_VI))};
const DEMOGRAPHICS = {_js([[t, list(c), req] for t, c, req in DEMOGRAPHICS_VI])};

function buildOne(f) {{
  const form = FormApp.create("NóiLái - phiếu " + ("0" + f.form).slice(-2));
  form.setCollectEmail(false);
  form.setDescription(INTRO);
  form.setProgressBar(true);
  const consent = form.addCheckboxItem().setTitle("Đồng ý tham gia").setChoiceValues(CONSENT).setRequired(true);
  consent.setValidation(FormApp.createCheckboxValidation().requireSelectExactly(CONSENT.length).build());
  form.addPageBreakItem().setTitle("Thông tin chung (thô, không định danh)");
  DEMOGRAPHICS.forEach(function (q) {{
    form.addMultipleChoiceItem().setTitle(q[0]).setChoiceValues(q[1]).setRequired(q[2]);
  }});
  form.addPageBreakItem().setTitle("Các câu hỏi");
  f.items.forEach(function (it) {{
    let item;
    if (it.yesno) {{
      item = form.addMultipleChoiceItem().setTitle(it.title).setChoiceValues(["Có", "Không"]);
    }} else {{
      item = form.addTextItem().setTitle(it.title);
    }}
    if (it.help) item.setHelpText(it.help);
    item.setRequired(it.required);
  }});
  form.addParagraphTextItem().setTitle("Góp ý (không bắt buộc)");
  return form.getPublishedUrl();
}}

function buildAll() {{
  const lines = [];
  FORMS.forEach(function (f) {{ lines.push(f.form + "\\t" + buildOne(f)); }});
  Logger.log(lines.join("\\n"));
}}
"""
    out = d / "build_forms.gs"
    out.write_text(js, encoding="utf-8")
    print(json.dumps({"written": str(out), "forms": len(forms), "items_per_form": [len(f["items"]) for f in forms][:3]}, ensure_ascii=False))
    return 0


def cmd_import_responses(args) -> int:
    """Google Forms response CSVs (responses_form_<nn>.csv) -> returned/baseline_form_<nn>_r<k>.csv in the form's own
    layout (answers by position), demographics to returned/demographics.csv (coarse bands only, no timestamps kept)."""
    d = Path(args.dir)
    ret = d / "returned"
    ret.mkdir(parents=True, exist_ok=True)
    demo_rows = []
    n = 0
    for path in sorted((d / "responses").glob("responses_form_*.csv")):
        nn = int(path.stem.rsplit("_", 1)[1])
        with open(d / f"baseline_form_{nn:02d}.csv", encoding="utf-8") as fh:
            form_rows = list(csv.DictReader(fh))
        by_pos = {int(r["position"]): r for r in form_rows}
        with open(path, encoding="utf-8") as fh:
            responses = list(csv.DictReader(fh))
        for k, resp in enumerate(responses, start=1):
            out_rows = []
            for col, val in resp.items():
                head = (col or "").split(".", 1)[0].strip()
                if head.isdigit() and int(head) in by_pos:
                    r = dict(by_pos[int(head)])
                    r["answer"] = val or ""
                    out_rows.append(r)
            out_rows.sort(key=lambda r: int(r["position"]))
            _write_csv(ret / f"baseline_form_{nn:02d}_r{k}.csv", list(form_rows[0].keys()), out_rows)
            demo_rows.append({"form": nn, "response": k, **{t: resp.get(t, "") for t, _c, _r in DEMOGRAPHICS_VI}})
            n += 1
    _write_csv(ret / "demographics.csv", ["form", "response", *[t for t, _c, _r in DEMOGRAPHICS_VI]], demo_rows)
    print(json.dumps({"responses": n, "written_to": str(ret)}, ensure_ascii=False))
    return 0


# --------------------------------------------------------------------------- cli
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("packet", help="Parts A-E for every validator + the author's keys")
    s.add_argument("--release", required=True)
    s.add_argument("--out", required=True)
    s.add_argument("--validators", nargs="+", default=["A", "B", "C"])
    s.add_argument("--attested", default=str(ROOT / "data" / "attested_seed.tsv"))
    s.add_argument("--candidates", default=str(ROOT / "data" / "attested_candidates.tsv"))
    s.add_argument("--per-cell", type=int, default=C.VALIDATION_PER_CELL)
    s.add_argument("--controls-per-cell", type=int, default=C.VALIDATION_CONTROLS_PER_CELL)
    s.add_argument("--overlap", type=int, default=C.VALIDATION_OVERLAP)
    s.add_argument("--t2-overlap", type=int, default=C.VALIDATION_T2_GOLD_OVERLAP)
    s.add_argument("--no-t2-gold", action="store_true", help="omit Part E")
    s.add_argument("--seed", type=int, default=PACKET_SEED)
    s.set_defaults(func=cmd_packet)
    a = sub.add_parser("adjudication-sheet", help="the split rows for the third validator, blind")
    a.add_argument("--dir", required=True)
    a.add_argument("--to", required=True)
    a.add_argument("--returned", nargs="+", default=None)
    a.add_argument("--fast", action="store_true")
    a.set_defaults(func=cmd_adjudication_sheet)
    c = sub.add_parser("score")
    c.add_argument("--dir", required=True)
    c.add_argument("--returned", nargs="+", default=None)
    c.add_argument("--n-boot", type=int, default=1000)
    c.set_defaults(func=cmd_score)
    b = sub.add_parser("baseline")
    b.add_argument("--items", required=True)
    b.add_argument("--out", required=True)
    b.add_argument("--attested", default=None, help="the release attested.jsonl, for the natural-competence block")
    b.add_argument("--n-forms", type=int, default=C.HUMAN_BASELINE_PEOPLE)
    b.add_argument("--per-form", type=int, default=C.HUMAN_BASELINE_ITEMS_PER_PERSON)
    b.add_argument("--anchors", type=int, default=C.HUMAN_BASELINE_ANCHORS)
    b.add_argument("--per-cell", type=int, default=None, help="double-judged items per task x variant cell (derived when omitted)")
    b.add_argument("--seed", type=int, default=BASELINE_SEED)
    b.add_argument("--no-model-prompt", action="store_true", help="omit the rendered model prompt (tests on synthetic items)")
    b.set_defaults(func=cmd_baseline)
    sb = sub.add_parser("score-baseline")
    sb.add_argument("--items", required=True)
    sb.add_argument("--dir", required=True)
    sb.add_argument("--returned", nargs="+", required=True)
    sb.add_argument("--min-answered", type=int, default=15)
    sb.set_defaults(func=cmd_score_baseline)
    g = sub.add_parser("google-form", help="Apps Script that builds the baseline Google Forms in the author's account")
    g.add_argument("--dir", required=True)
    g.set_defaults(func=cmd_google_form)
    ir = sub.add_parser("import-responses", help="Google Forms response CSVs -> returned/ in the forms' layout")
    ir.add_argument("--dir", required=True)
    ir.set_defaults(func=cmd_import_responses)
    args = ap.parse_args(argv)
    if getattr(args, "returned", None) is None and args.cmd in ("score", "adjudication-sheet"):
        args.returned = [str(Path(args.dir) / "returned" / "*")]
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
